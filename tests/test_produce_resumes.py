from __future__ import annotations

import base64
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from scripts.mark_applied_jobs import TrackerJob, write_tracker_jobs_snapshot
from scripts.produce_resumes import (
    SourceTemplate,
    ResumeProductionResult,
    build_resume_from_template,
    docker_renderer_image,
    load_template_specs,
    main,
    preflight_sources,
    produce_resumes,
    render_docx,
    render_docx_in_docker,
    render_docx_native,
    resolve_binary,
    validate_docx,
    validate_rendered_pages,
)
from scripts.resume_artifacts import parse_resume_ready_jobs, resume_artifact_name, resume_artifact_names


ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def write_source_template(path: Path, *, title: str = "Senior Principal Software Engineer") -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    left = table.cell(0, 0)
    left.paragraphs[0].add_run("Candidate Name")
    left.add_paragraph(title)
    table.cell(0, 1).paragraphs[0].add_run("candidate@example.com")
    document.add_paragraph(
        "Senior principal software engineer with over 25 years of experience building enterprise software."
    )
    document.add_paragraph("Areas of Expertise")
    document.save(path)


def software_template_specs(source_path: Path) -> dict[str, SourceTemplate]:
    return {
        "software_engineering": SourceTemplate(
            filename=source_path.name,
            sha256=hashlib.sha256(source_path.read_bytes()).hexdigest(),
            source_title="Senior Principal Software Engineer",
            summary_prefix="Senior principal software engineer",
        )
    }


def write_ready_queue(path: Path, rows: list[str], *, run_date: str = "2026-07-22") -> None:
    path.write_text(
        "\n".join(
            [
                f"# {run_date} - Resume-Ready Jobs",
                "",
                "Only jobs in this file should proceed to Word resume or cover-letter generation.",
                "",
                "| Priority | Company | Role | Decision | Template | Fit | URL |",
                "|---:|---|---|---|---|---:|---|",
                *rows,
                "",
            ]
        ),
        encoding="utf-8",
    )


def write_workflow_review(
    path: Path,
    *,
    company: str,
    title: str,
    run_date: str = "2026-07-22",
    url: str = "https://example.com/job",
) -> None:
    payload = {
        "generated_on": run_date,
        "reviews": [
            {
                "company": company,
                "title": title,
                "url": url,
                "workflow_decision": "APPLY",
                "should_generate_resume": True,
                "source_fit_score": 86,
            }
        ],
    }
    path.write_text(json.dumps(payload), encoding="utf-8")


def fake_renderer(document_path: Path, render_dir: Path) -> list[Path]:
    if not document_path.exists():
        raise FileNotFoundError(document_path)
    render_dir.mkdir(parents=True, exist_ok=True)
    page = render_dir / "page-1.png"
    page.write_bytes(ONE_PIXEL_PNG)
    return [page]


def write_empty_applied_snapshot(root: Path) -> Path:
    path = root / "2026-07-22-applied-jobs.json"
    write_tracker_jobs_snapshot(path, jobs=[], ssh_host="test-host")
    return path


class ProduceResumesTests(unittest.TestCase):
    def setUp(self) -> None:
        contact_environment = patch.dict(
            os.environ,
            {
                "JOBAGENT_CANDIDATE_EMAIL": "candidate@example.com",
                "JOBAGENT_RESUME_LINKEDIN_URL": "https://www.linkedin.com/in/candidate-profile",
                "JOBAGENT_RESUME_PORTFOLIO_URL": "https://portfolio.example.com",
            },
        )
        contact_environment.start()
        self.addCleanup(contact_environment.stop)

    def test_contact_links_read_runtime_candidate_configuration(self) -> None:
        from scripts.produce_resumes import configured_contact_links

        with patch.dict(os.environ, {
            "JOBAGENT_CANDIDATE_EMAIL": "engineer@example.org",
            "JOBAGENT_RESUME_LINKEDIN_URL": "https://www.linkedin.com/in/example-engineer",
            "JOBAGENT_RESUME_PORTFOLIO_URL": "https://work.example.org",
        }):
            self.assertEqual(
                (("Email", "mailto:engineer@example.org"),
                 ("LinkedIn", "https://www.linkedin.com/in/example-engineer"),
                 ("AI Portfolio", "https://work.example.org")),
                configured_contact_links(),
            )

    def test_invalid_contact_configuration_is_rejected(self) -> None:
        from scripts.produce_resumes import configured_contact_links

        for key, value in (("JOBAGENT_CANDIDATE_EMAIL", ""),
                           ("JOBAGENT_CANDIDATE_EMAIL", "invalid"),
                           ("JOBAGENT_RESUME_LINKEDIN_URL", "javascript:alert(1)"),
                           ("JOBAGENT_RESUME_PORTFOLIO_URL", "https://")):
            with self.subTest(key=key, value=value), patch.dict(os.environ, {key: value}):
                with self.assertRaises(ValueError):
                    configured_contact_links()

    def test_resume_artifact_name_uses_configured_candidate_name(self) -> None:
        name = resume_artifact_name(
            "2026-07-24",
            "ReadyCo",
            "Principal Engineer",
            candidate_name="Candidate Name",
        )

        self.assertEqual(
            "2026-07-24-ReadyCo-Principal-Engineer-Candidate-Name.docx",
            name,
        )

    def test_generated_resume_preserves_word_compatibility_namespaces(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_path = root / "source.docx"
            write_source_template(source_path)
            with zipfile.ZipFile(source_path, "r") as source:
                parts = {entry.filename: source.read(entry.filename) for entry in source.infolist()}
            document_xml = parts["word/document.xml"]
            document_xml = document_xml.replace(
                b"<w:document ",
                b'<w:document xmlns:acme="urn:jobagent:word-compat" ',
                1,
            )
            document_xml = document_xml.replace(
                b'mc:Ignorable="w14 wp14"',
                b'mc:Ignorable="w14 wp14 acme"',
                1,
            )
            parts["word/document.xml"] = document_xml
            with zipfile.ZipFile(source_path, "w") as destination:
                for name, content in parts.items():
                    destination.writestr(name, content)

            output_path = root / "resume.docx"
            build_resume_from_template(
                source_path=source_path,
                destination_path=output_path,
                job=parse_resume_ready_jobs(
                    self._write_single_ready_queue(root, "software_engineering")
                )[0],
                spec=software_template_specs(source_path)["software_engineering"],
            )

            with zipfile.ZipFile(output_path, "r") as generated:
                generated_xml = generated.read("word/document.xml")
                settings_root = ET.fromstring(generated.read("word/settings.xml"))
                document_root = ET.fromstring(generated_xml)
                relationships_root = ET.fromstring(
                    generated.read("word/_rels/document.xml.rels")
                )

            compatibility_mode = next(
                setting
                for setting in settings_root.findall(
                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}compatSetting"
                )
                if setting.attrib.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name"
                )
                == "compatibilityMode"
            )
            opening_summary = next(
                paragraph
                for paragraph in document_root.findall(
                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
                )
                if "".join(
                    node.text or ""
                    for node in paragraph.findall(
                        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                    )
                )
                .startswith("Principal Engineer with over 25 years")
            )
            summary_spacing = opening_summary.find(
                "./{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr/"
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing"
            )
            contact_paragraph = next(
                paragraph
                for paragraph in document_root.findall(
                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
                )
                if "".join(
                    node.text or ""
                    for node in paragraph.findall(
                        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                    )
                )
                == "Email • LinkedIn • AI Portfolio"
            )
            relationship_targets = {
                relationship.attrib["Id"]: relationship.attrib["Target"]
                for relationship in relationships_root.findall(
                    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
                )
            }
            contact_targets = {
                "".join(
                    node.text or ""
                    for node in hyperlink.findall(
                        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                    )
                ): relationship_targets[
                    hyperlink.attrib[
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    ]
                ]
                for hyperlink in contact_paragraph.findall(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
                )
            }

        self.assertIn(b'xmlns:acme="urn:jobagent:word-compat"', generated_xml)
        self.assertIn(b'mc:Ignorable="w14 wp14 acme"', generated_xml)
        self.assertEqual(
            "15",
            compatibility_mode.attrib[
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
            ],
        )
        self.assertIsNotNone(summary_spacing)
        self.assertEqual(
            "240",
            summary_spacing.attrib[
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}before"
            ]
            if summary_spacing is not None
            else "",
        )
        self.assertEqual(
            {
                "Email": "mailto:candidate@example.com",
                "LinkedIn": "https://www.linkedin.com/in/candidate-profile",
                "AI Portfolio": "https://portfolio.example.com",
            },
            contact_targets,
        )

    def test_senior_execution_positioning_softens_principal_summary_signal(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_path = root / "source.docx"
            write_source_template(source_path)
            ready_path = root / "2026-07-22-resume-ready-jobs.md"
            ready_path.write_text(
                "\n".join(
                    [
                        "# 2026-07-22 - Resume-Ready Jobs",
                        "",
                        "| Priority | Company | Role | Decision | Template | Fit | Overqualified Risk | Resume Positioning | URL |",
                        "|---:|---|---|---|---|---:|---|---|---|",
                        "| 1 | ReadyCo | Senior Software Engineer | APPLY | software_engineering | 86 | medium | senior_lead_execution_forward | https://example.com/job |",
                    ]
                ),
                encoding="utf-8",
            )
            output_path = root / "resume.docx"

            build_resume_from_template(
                source_path=source_path,
                destination_path=output_path,
                job=parse_resume_ready_jobs(ready_path)[0],
                spec=software_template_specs(source_path)["software_engineering"],
            )

            rendered_text = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)

        self.assertIn("Senior Software Engineer with extensive experience", rendered_text)
        self.assertNotIn("over 25 years", rendered_text)

    def test_validation_rejects_legacy_word_compatibility_mode(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "source.docx"
            write_source_template(source_path)

            with self.assertRaisesRegex(ValueError, "compatibility level 15"):
                validate_docx(
                    source_path=source_path,
                    resume_path=source_path,
                    expected_title="Senior Principal Software Engineer",
                )

    def test_main_produces_resumes_and_cover_letters_for_ready_queue(self) -> None:
        result = ResumeProductionResult(
            generated_on="2026-07-22",
            records=[],
            manifest_path=Path("outputs/2026-07-22-resume-production.json"),
            workflow_html_path=Path("outputs/2026-07-22-workflow-review.html"),
        )
        with (
            patch.object(
                sys,
                "argv",
                [
                    "produce_resumes.py",
                    "--resume-ready",
                    "outputs/2026-07-22-resume-ready-jobs.md",
                    "--applied-jobs",
                    "outputs/2026-07-22-applied-jobs.json",
                ],
            ),
            patch("scripts.produce_resumes.produce_resumes", return_value=result) as producer,
            patch(
                "scripts.produce_resumes.produce_cover_letters_for_ready_queue",
                return_value=SimpleNamespace(
                    records=[
                        SimpleNamespace(status="generated"),
                        SimpleNamespace(status="existing"),
                    ],
                    manifest_path=Path("outputs/2026-07-22-cover-letter-production.json"),
                ),
            ) as cover_letter_producer,
            patch("scripts.produce_resumes.write_workflow_review_html", return_value=result.workflow_html_path),
            patch("builtins.print") as printed,
        ):
            main()

        self.assertTrue(producer.called)
        cover_letter_producer.assert_called_once_with(
            resume_ready_path=Path("outputs/2026-07-22-resume-ready-jobs.md"),
            applied_jobs_path=Path("outputs/2026-07-22-applied-jobs.json"),
            source_dir=Path("Source Docs"),
            resume_output_dir=Path("outputs/resumes"),
        )
        self.assertEqual(5, printed.call_count)
        printed.assert_any_call("Validated 2 cover letter(s): 1 generated, 1 existing")
        printed.assert_any_call(
            "Wrote cover-letter production manifest: "
            "outputs/2026-07-22-cover-letter-production.json"
        )

    def test_resolve_binary_prefers_path_then_macos_fallback(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            fallback = Path(tmp_dir) / "fallback"
            fallback.write_text("binary", encoding="utf-8")
            with patch("scripts.produce_resumes.shutil.which", return_value="/usr/bin/example"):
                self.assertEqual("/usr/bin/example", resolve_binary("example", fallback))
            with patch("scripts.produce_resumes.shutil.which", return_value=None):
                self.assertEqual(str(fallback), resolve_binary("example", fallback))
            with patch("scripts.produce_resumes.shutil.which", return_value=None):
                with self.assertRaisesRegex(FileNotFoundError, "unavailable"):
                    resolve_binary("example")

    def test_template_configuration_and_preflight_validation_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            config_path = root / "templates.json"
            config_path.write_text(
                json.dumps(
                    {
                        "software_engineering": {
                            "filename": "source.docx",
                            "sha256": "a" * 64,
                            "source_title": "Engineer",
                            "summary_prefix": "Engineer",
                        }
                    }
                ),
                encoding="utf-8",
            )
            specs = load_template_specs(config_path)
            jobs = parse_resume_ready_jobs(self._write_single_ready_queue(root, "software_engineering"))
            self.assertEqual("source.docx", specs["software_engineering"].filename)
            with self.assertRaisesRegex(ValueError, "No source resume mapping"):
                preflight_sources(jobs, root, {})
            with self.assertRaisesRegex(FileNotFoundError, "Missing source resume"):
                preflight_sources(jobs, root, specs)

    def test_resume_ready_parser_preserves_seniority_risk_and_positioning(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            ready_path = Path(tmp_dir) / "2026-07-22-resume-ready-jobs.md"
            ready_path.write_text(
                "\n".join(
                    [
                        "# 2026-07-22 - Resume-Ready Jobs",
                        "",
                        "| Priority | Company | Role | Decision | Template | Fit | Overqualified Risk | Resume Positioning | URL |",
                        "|---:|---|---|---|---|---:|---|---|---|",
                        "| 1 | ReadyCo | Lead Software Engineer | APPLY | software_engineering | 86 | medium | senior_lead_execution_forward | https://example.com/job |",
                    ]
                ),
                encoding="utf-8",
            )

            job = parse_resume_ready_jobs(ready_path)[0]

        self.assertEqual("medium", job.overqualified_risk)
        self.assertEqual("senior_lead_execution_forward", job.resume_positioning)

    @staticmethod
    def _write_single_ready_queue(root: Path, template: str) -> Path:
        ready_path = root / "2026-07-22-resume-ready-jobs.md"
        write_ready_queue(
            ready_path,
            [f"| 1 | ReadyCo | Principal Engineer | APPLY | {template} | 86 | https://example.com/job |"],
        )
        return ready_path

    def test_document_and_page_validation_rejects_invalid_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_path = root / "source.docx"
            write_source_template(source_path)
            invalid_path = root / "invalid.docx"
            invalid_path.write_text("not a docx", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "not a valid DOCX"):
                validate_docx(source_path=source_path, resume_path=invalid_path, expected_title="Engineer")
            with self.assertRaisesRegex(ValueError, "missing or empty"):
                validate_docx(source_path=source_path, resume_path=root / "missing.docx", expected_title="Engineer")
            with self.assertRaisesRegex(ValueError, "produced no pages"):
                validate_rendered_pages([], document_path=source_path)
            invalid_page = root / "page-1.png"
            invalid_page.write_bytes(b"not a png")
            with self.assertRaisesRegex(ValueError, "invalid PNG"):
                validate_rendered_pages([invalid_page], document_path=source_path)

    def test_native_renderer_returns_rasterized_pages(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            document_path = root / "resume.docx"
            document_path.write_bytes(b"placeholder")
            render_dir = root / "render"

            def fake_run(command: list[str], **_kwargs: object) -> subprocess.CompletedProcess[str]:
                if command[0] == "soffice":
                    render_dir.mkdir(parents=True, exist_ok=True)
                    (render_dir / "resume.pdf").write_bytes(b"%PDF")
                    return subprocess.CompletedProcess(command, 0, "", "")
                if command[0] == "pdftoppm":
                    (render_dir / "page-1.png").write_bytes(ONE_PIXEL_PNG)
                    return subprocess.CompletedProcess(command, 0, "", "")
                raise AssertionError(f"Unexpected renderer command: {command}")

            with (
                patch(
                    "scripts.produce_resumes.resolve_binary",
                    side_effect=lambda name, _fallback=None: name,
                ),
                patch("scripts.produce_resumes.subprocess.run", side_effect=fake_run),
            ):
                pages = render_docx_native(document_path, render_dir)

        self.assertEqual(["page-1.png"], [page.name for page in pages])

    def test_docker_renderer_requires_pages_and_nonempty_image_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            document_path = root / "resume.docx"
            document_path.write_bytes(b"placeholder")
            with (
                patch("scripts.produce_resumes.resolve_binary", return_value="docker"),
                patch(
                    "scripts.produce_resumes.subprocess.run",
                    return_value=subprocess.CompletedProcess(["docker"], 0, "", ""),
                ),
                self.assertRaisesRegex(RuntimeError, "did not create rendered pages"),
            ):
                render_docx_in_docker(document_path, root / "render")
            with patch.dict(os.environ, {"JOBAGENT_DOCKER_RENDERER_IMAGE": "   "}):
                with self.assertRaisesRegex(ValueError, "cannot be empty"):
                    docker_renderer_image()

    def test_render_docx_falls_back_to_isolated_docker_renderer(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            document_path = root / "resume.docx"
            document_path.write_bytes(b"placeholder")
            render_dir = root / "render"
            commands: list[list[str]] = []

            def fake_run(command: list[str], **_kwargs: object) -> subprocess.CompletedProcess[str]:
                commands.append(command)
                if command[0] == "soffice":
                    return subprocess.CompletedProcess(command, 6, "", "native Abort trap")
                if command[0] == "docker":
                    render_dir.mkdir(parents=True, exist_ok=True)
                    (render_dir / "page-1.png").write_bytes(ONE_PIXEL_PNG)
                    return subprocess.CompletedProcess(command, 0, "", "")
                raise AssertionError(f"Unexpected renderer command: {command}")

            with (
                patch(
                    "scripts.produce_resumes.resolve_binary",
                    side_effect=lambda name, _fallback=None: name,
                ),
                patch("scripts.produce_resumes.subprocess.run", side_effect=fake_run),
            ):
                pages = render_docx(document_path, render_dir)

        self.assertEqual(["page-1.png"], [page.name for page in pages])
        self.assertEqual("soffice", commands[0][0])
        docker_command = next(command for command in commands if command[0] == "docker")
        self.assertIn("--network", docker_command)
        self.assertIn("none", docker_command)
        self.assertIn("jobsearchagent-libreoffice-renderer:1.0.0", docker_command)

    def test_render_docx_falls_back_when_native_renderer_times_out(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            document_path = root / "resume.docx"
            document_path.write_bytes(b"placeholder")
            render_dir = root / "render"

            def fake_run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[str]:
                self.assertIn("timeout", kwargs)
                if command[0] == "soffice":
                    raise subprocess.TimeoutExpired(command, kwargs["timeout"])
                if command[0] == "docker":
                    render_dir.mkdir(parents=True, exist_ok=True)
                    (render_dir / "page-1.png").write_bytes(ONE_PIXEL_PNG)
                    return subprocess.CompletedProcess(command, 0, "", "")
                raise AssertionError(f"Unexpected renderer command: {command}")

            with (
                patch(
                    "scripts.produce_resumes.resolve_binary",
                    side_effect=lambda name, _fallback=None: name,
                ),
                patch("scripts.produce_resumes.subprocess.run", side_effect=fake_run),
            ):
                pages = render_docx(document_path, render_dir)

        self.assertEqual(["page-1.png"], [page.name for page in pages])

    def test_render_docx_can_use_docker_without_trying_native_renderer(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            document_path = root / "resume.docx"
            document_path.write_bytes(b"placeholder")
            render_dir = root / "render"
            commands: list[list[str]] = []

            def fake_run(command: list[str], **_kwargs: object) -> subprocess.CompletedProcess[str]:
                commands.append(command)
                render_dir.mkdir(parents=True, exist_ok=True)
                (render_dir / "page-1.png").write_bytes(ONE_PIXEL_PNG)
                return subprocess.CompletedProcess(command, 0, "", "")

            with (
                patch.dict(os.environ, {"JOBAGENT_RENDERER_MODE": "docker"}),
                patch("scripts.produce_resumes.resolve_binary", return_value="docker"),
                patch("scripts.produce_resumes.subprocess.run", side_effect=fake_run),
            ):
                pages = render_docx(document_path, render_dir)

        self.assertEqual(["page-1.png"], [page.name for page in pages])
        self.assertEqual(1, len(commands))
        self.assertEqual("docker", commands[0][0])
        self.assertFalse(any(command[0] == "soffice" for command in commands))

    def test_render_docx_reports_both_native_and_docker_failures(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            document_path = root / "resume.docx"
            document_path.write_bytes(b"placeholder")

            def fake_run(command: list[str], **_kwargs: object) -> subprocess.CompletedProcess[str]:
                if command[0] == "soffice":
                    return subprocess.CompletedProcess(command, 6, "", "native Abort trap")
                if command[0] == "docker":
                    return subprocess.CompletedProcess(command, 125, "", "container image unavailable")
                raise AssertionError(f"Unexpected renderer command: {command}")

            with (
                patch(
                    "scripts.produce_resumes.resolve_binary",
                    side_effect=lambda name, _fallback=None: name,
                ),
                patch("scripts.produce_resumes.subprocess.run", side_effect=fake_run),
                self.assertRaisesRegex(RuntimeError, "native renderer failed.*container renderer failed"),
            ):
                render_docx(document_path, root / "render")

    def test_same_run_queue_keeps_distinct_posting_urls_with_same_company_and_title(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            ready_path = Path(tmp_dir) / "2026-07-24-resume-ready-jobs.md"
            write_ready_queue(
                ready_path,
                [
                    "| 1 | Block | Senior Software Engineer, Cash App Taxes | APPLY | software_engineering | 79 | http://block.xyz/careers/jobs/5196187008?gh_jid=5196187008 |",
                    "| 2 | Block | Senior Software Engineer, Cash App Taxes | APPLY | software_engineering | 79 | http://block.xyz/careers/jobs/5207134008?gh_jid=5207134008 |",
                ],
                run_date="2026-07-24",
            )

            jobs = parse_resume_ready_jobs(ready_path)
            names = resume_artifact_names("2026-07-24", jobs)

        self.assertEqual(2, len(jobs))
        self.assertEqual(2, len(set(names.values())))

    def test_generates_only_ready_rows_validates_and_regenerates_html(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            source_path = source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx"
            write_source_template(source_path)
            ready_path = root / "2026-07-22-resume-ready-jobs.md"
            write_ready_queue(
                ready_path,
                ["| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | https://example.com/job |"],
            )
            write_workflow_review(
                root / "2026-07-22-workflow-review.json",
                company="ReadyCo",
                title="Principal Engineer",
            )

            result = produce_resumes(
                resume_ready_path=ready_path,
                applied_jobs_path=write_empty_applied_snapshot(root),
                source_dir=source_dir,
                output_dir=root / "resumes",
                render_fn=fake_renderer,
                template_specs=software_template_specs(source_path),
            )

            resume_path = root / "resumes" / "2026-07-22-ReadyCo-Principal-Engineer-Candidate-Name.docx"
            self.assertTrue(resume_path.exists())
            self.assertEqual("Principal Engineer", Document(resume_path).tables[0].cell(0, 0).paragraphs[1].text)
            self.assertTrue(Document(resume_path).paragraphs[0].text.startswith("Principal Engineer with over 25 years"))
            self.assertEqual("generated", result.records[0].status)
            self.assertEqual(1, result.records[0].page_count)
            manifest = json.loads((root / "2026-07-22-resume-production.json").read_text(encoding="utf-8"))
            self.assertEqual("validated", manifest["records"][0]["validation_status"])
            detail_path = next((root / "2026-07-22-workflow-review-details").glob("*.html"))
            detail_html = detail_path.read_text(encoding="utf-8")
            self.assertIn("Open generated resume", detail_html)
            self.assertIn("../resumes/2026-07-22-ReadyCo-Principal-Engineer-Candidate-Name.docx", detail_html)

    def test_second_run_is_idempotent_and_revalidates_existing_resume(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            source_path = source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx"
            write_source_template(source_path)
            ready_path = root / "2026-07-22-resume-ready-jobs.md"
            write_ready_queue(
                ready_path,
                ["| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | https://example.com/job |"],
            )
            write_workflow_review(
                root / "2026-07-22-workflow-review.json",
                company="ReadyCo",
                title="Principal Engineer",
                url="https://example.com/jobs/123456",
            )
            output_dir = root / "resumes"

            produce_resumes(
                resume_ready_path=ready_path,
                applied_jobs_path=write_empty_applied_snapshot(root),
                source_dir=source_dir,
                output_dir=output_dir,
                render_fn=fake_renderer,
                template_specs=software_template_specs(source_path),
            )
            resume_path = output_dir / "2026-07-22-ReadyCo-Principal-Engineer-Candidate-Name.docx"
            first_hash = hashlib.sha256(resume_path.read_bytes()).hexdigest()
            with zipfile.ZipFile(resume_path, "r") as resume_archive:
                parts = {
                    entry.filename: resume_archive.read(entry.filename)
                    for entry in resume_archive.infolist()
                }
            parts["word/document.xml"] = ET.tostring(
                ET.fromstring(parts["word/document.xml"]),
                encoding="utf-8",
                xml_declaration=True,
            )
            with zipfile.ZipFile(resume_path, "w") as resume_archive:
                for name, content in parts.items():
                    resume_archive.writestr(name, content)
            corrupt_hash = hashlib.sha256(resume_path.read_bytes()).hexdigest()

            result = produce_resumes(
                resume_ready_path=ready_path,
                applied_jobs_path=write_empty_applied_snapshot(root),
                source_dir=source_dir,
                output_dir=output_dir,
                render_fn=fake_renderer,
                template_specs=software_template_specs(source_path),
            )

            self.assertNotEqual(corrupt_hash, hashlib.sha256(resume_path.read_bytes()).hexdigest())
            self.assertEqual(first_hash, hashlib.sha256(resume_path.read_bytes()).hexdigest())
            self.assertEqual("generated", result.records[0].status)
            self.assertEqual("validated", result.records[0].validation_status)

    def test_new_day_reuses_exact_job_resume_and_links_prior_artifact(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            source_path = source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx"
            write_source_template(source_path)
            output_dir = root / "resumes"
            ready_row = (
                "| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | "
                "https://example.com/jobs/123456 |"
            )

            first_ready_path = root / "2026-07-22-resume-ready-jobs.md"
            write_ready_queue(first_ready_path, [ready_row])
            write_workflow_review(
                root / "2026-07-22-workflow-review.json",
                company="ReadyCo",
                title="Principal Engineer",
                url="https://example.com/jobs/123456",
            )
            produce_resumes(
                resume_ready_path=first_ready_path,
                applied_jobs_path=write_empty_applied_snapshot(root),
                source_dir=source_dir,
                output_dir=output_dir,
                render_fn=fake_renderer,
                template_specs=software_template_specs(source_path),
            )
            prior_path = output_dir / "2026-07-22-ReadyCo-Principal-Engineer-Candidate-Name.docx"
            with zipfile.ZipFile(prior_path, "r") as resume_archive:
                parts = {
                    entry.filename: resume_archive.read(entry.filename)
                    for entry in resume_archive.infolist()
                }
            parts["word/document.xml"] = ET.tostring(
                ET.fromstring(parts["word/document.xml"]),
                encoding="utf-8",
                xml_declaration=True,
            )
            with zipfile.ZipFile(prior_path, "w") as resume_archive:
                for name, content in parts.items():
                    resume_archive.writestr(name, content)

            second_ready_path = root / "2026-07-23-resume-ready-jobs.md"
            write_ready_queue(second_ready_path, [ready_row], run_date="2026-07-23")
            write_workflow_review(
                root / "2026-07-23-workflow-review.json",
                company="ReadyCo",
                title="Principal Engineer",
                run_date="2026-07-23",
                url="https://example.com/jobs/123456",
            )
            second_applied_path = root / "2026-07-23-applied-jobs.json"
            write_tracker_jobs_snapshot(second_applied_path, jobs=[], ssh_host="test-host")

            result = produce_resumes(
                resume_ready_path=second_ready_path,
                applied_jobs_path=second_applied_path,
                source_dir=source_dir,
                output_dir=output_dir,
                render_fn=fake_renderer,
                template_specs=software_template_specs(source_path),
            )

            prior_name = "2026-07-22-ReadyCo-Principal-Engineer-Candidate-Name.docx"
            self.assertEqual([prior_name], [path.name for path in output_dir.glob("*.docx")])
            self.assertEqual("generated", result.records[0].status)
            self.assertEqual(str(output_dir / prior_name), result.records[0].artifact_path)
            detail_path = next((root / "2026-07-23-workflow-review-details").glob("*.html"))
            detail_html = detail_path.read_text(encoding="utf-8")
            self.assertIn(f'href="../resumes/{prior_name}"', detail_html)
            self.assertIn("Open generated resume", detail_html)

    def test_same_company_and_title_with_different_job_id_generates_separate_resume(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            source_path = source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx"
            write_source_template(source_path)
            output_dir = root / "resumes"
            first_url = "https://example.com/jobs/123456"
            second_url = "https://example.com/jobs/654321"

            first_ready_path = root / "2026-07-22-resume-ready-jobs.md"
            write_ready_queue(
                first_ready_path,
                [f"| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | {first_url} |"],
            )
            write_workflow_review(
                root / "2026-07-22-workflow-review.json",
                company="ReadyCo",
                title="Principal Engineer",
                url=first_url,
            )
            produce_resumes(
                resume_ready_path=first_ready_path,
                applied_jobs_path=write_empty_applied_snapshot(root),
                source_dir=source_dir,
                output_dir=output_dir,
                render_fn=fake_renderer,
                template_specs=software_template_specs(source_path),
            )

            second_ready_path = root / "2026-07-23-resume-ready-jobs.md"
            write_ready_queue(
                second_ready_path,
                [f"| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | {second_url} |"],
                run_date="2026-07-23",
            )
            write_workflow_review(
                root / "2026-07-23-workflow-review.json",
                company="ReadyCo",
                title="Principal Engineer",
                run_date="2026-07-23",
                url=second_url,
            )
            second_applied_path = root / "2026-07-23-applied-jobs.json"
            write_tracker_jobs_snapshot(second_applied_path, jobs=[], ssh_host="test-host")

            result = produce_resumes(
                resume_ready_path=second_ready_path,
                applied_jobs_path=second_applied_path,
                source_dir=source_dir,
                output_dir=output_dir,
                render_fn=fake_renderer,
                template_specs=software_template_specs(source_path),
            )

            self.assertEqual("generated", result.records[0].status)
            self.assertEqual(2, len(list(output_dir.glob("*.docx"))))

    def test_rejects_non_apply_decision_before_writing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            ready_path = root / "2026-07-22-resume-ready-jobs.md"
            write_ready_queue(
                ready_path,
                ["| 1 | StretchCo | Engineer | MODERATE STRETCH | software_engineering | 68 | https://example.com/stretch |"],
            )

            with self.assertRaisesRegex(ValueError, "not authorized"):
                produce_resumes(
                    resume_ready_path=ready_path,
                    applied_jobs_path=write_empty_applied_snapshot(root),
                    source_dir=root / "Source Docs",
                    output_dir=root / "resumes",
                    render_fn=fake_renderer,
                    template_specs={},
                )

            self.assertFalse((root / "resumes").exists())

    def test_missing_source_fails_preflight_without_partial_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            source_path = source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx"
            write_source_template(source_path)
            ready_path = root / "2026-07-22-resume-ready-jobs.md"
            write_ready_queue(
                ready_path,
                [
                    "| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | https://example.com/one |",
                    "| 2 | MissingCo | Solutions Engineer | APPLY | technical_success | 85 | https://example.com/two |",
                ],
            )

            with self.assertRaisesRegex(FileNotFoundError, "Technical_Success"):
                produce_resumes(
                    resume_ready_path=ready_path,
                    applied_jobs_path=write_empty_applied_snapshot(root),
                    source_dir=source_dir,
                    output_dir=root / "resumes",
                    render_fn=fake_renderer,
                    template_specs={
                        **software_template_specs(source_path),
                        "technical_success": SourceTemplate(
                            filename="Candidate_Name_SOURCE_Technical_Success_Architect.docx",
                            sha256="missing",
                            source_title="Technical Success Architect",
                            summary_prefix="Technical success architect and enterprise SaaS engineering leader",
                        ),
                    },
                )

            self.assertFalse((root / "resumes").exists())

    def test_render_failure_does_not_publish_new_resume_or_html(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            source_path = source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx"
            write_source_template(source_path)
            ready_path = root / "2026-07-22-resume-ready-jobs.md"
            write_ready_queue(
                ready_path,
                ["| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | https://example.com/job |"],
            )
            write_workflow_review(
                root / "2026-07-22-workflow-review.json",
                company="ReadyCo",
                title="Principal Engineer",
            )

            def failing_renderer(_document_path: Path, _render_dir: Path) -> list[Path]:
                raise RuntimeError("render failed")

            with self.assertRaisesRegex(RuntimeError, "render failed"):
                produce_resumes(
                    resume_ready_path=ready_path,
                    applied_jobs_path=write_empty_applied_snapshot(root),
                    source_dir=source_dir,
                    output_dir=root / "resumes",
                    render_fn=failing_renderer,
                    template_specs=software_template_specs(source_path),
                )

            self.assertFalse((root / "resumes" / "2026-07-22-ReadyCo-Principal-Engineer-Candidate-Name.docx").exists())
            self.assertFalse((root / "2026-07-22-workflow-review.html").exists())

    def test_applied_job_in_stale_ready_queue_fails_closed_before_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            ready_path = root / "2026-07-22-resume-ready-jobs.md"
            write_ready_queue(
                ready_path,
                ["| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | https://example.com/jobs/123456 |"],
            )
            applied_jobs_path = root / "2026-07-22-applied-jobs.json"
            write_tracker_jobs_snapshot(
                applied_jobs_path,
                jobs=[
                    TrackerJob(
                        company="ReadyCo",
                        title="Principal Engineer",
                        status="applied",
                        job_url="https://example.com/jobs/123456",
                        applied_date="2026-07-20",
                    )
                ],
                ssh_host="test-host",
            )

            with self.assertRaisesRegex(ValueError, "already applied"):
                produce_resumes(
                    resume_ready_path=ready_path,
                    applied_jobs_path=applied_jobs_path,
                    source_dir=root / "Source Docs",
                    output_dir=root / "resumes",
                    render_fn=fake_renderer,
                    template_specs={},
                )

            self.assertFalse((root / "resumes").exists())
