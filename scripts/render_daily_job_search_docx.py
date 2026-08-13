#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from scripts.render_daily_job_search_html import (
    IntroductionMessage,
    NetworkingCompanySection,
    ReportRow,
    build_default_paths,
    extract_title,
    parse_networking_markdown,
    parse_report_rows,
    remove_networking_appendix,
)


SUMMARY_TOP_BOOKMARK = "summary_top"
SUMMARY_BOOKMARKS = {
    "Apply": "summary_apply",
    "Maybe": "summary_maybe",
}
TABLE_FONT_SIZE = Pt(8)
BODY_FONT_SIZE = Pt(10)


@dataclass(frozen=True)
class ReportSections:
    preamble_lines: list[str]
    section_lines: dict[str, list[str]]


class BookmarkManager:
    def __init__(self) -> None:
        self._next_id = 1

    def add_paragraph_bookmark(self, paragraph: Paragraph, name: str) -> None:
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), str(self._next_id))
        bookmark_start.set(qn("w:name"), name)

        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), str(self._next_id))

        paragraph._p.insert(0, bookmark_start)
        paragraph._p.append(bookmark_end)
        self._next_id += 1


def build_default_paths_docx(report_path: Path) -> tuple[Path, Path]:
    networking_path, _html_output_path = build_default_paths(report_path)
    output_path = report_path.with_name(f"{report_path.stem}.docx")
    return networking_path, output_path


def build_row_bookmark_name(row: ReportRow) -> str:
    key = f"{row.decision}|{row.priority}|{row.company}|{row.role}|{row.posted}"
    digest = hashlib.sha1(key.encode("utf-8")).hexdigest()[:12]
    return f"job_{row.decision.lower()}_{row.priority}_{digest}"


def parse_report_sections(cleaned_markdown: str) -> ReportSections:
    lines = cleaned_markdown.splitlines()
    if not lines:
        raise ValueError("Report markdown is empty.")

    preamble_lines: list[str] = []
    section_lines: dict[str, list[str]] = {}
    current_section: str | None = None

    for line in lines[1:]:
        stripped = line.strip()
        heading_match = re.fullmatch(r"##\s+(.*)", stripped)
        if heading_match:
            current_section = heading_match.group(1).strip()
            section_lines[current_section] = []
            continue

        if current_section is None:
            preamble_lines.append(line)
        else:
            section_lines[current_section].append(line)

    return ReportSections(
        preamble_lines=preamble_lines,
        section_lines=section_lines,
    )


def normalize_inline_markdown(text: str) -> str:
    normalized = re.sub(r"`([^`]+)`", r"\1", text)
    normalized = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", normalized)
    return normalized


def set_run_font(run_text_container: object, *, bold: bool = False, size: Pt = BODY_FONT_SIZE) -> None:
    run_text_container.bold = bold
    run_text_container.font.size = size


def append_hyperlink_run(paragraph: Paragraph, *, text: str, anchor: str | None, url: str | None) -> None:
    if anchor is None and url is None:
        raise ValueError("A hyperlink must target either an internal anchor or an external URL.")
    if anchor is not None and url is not None:
        raise ValueError("A hyperlink cannot target both an internal anchor and an external URL.")

    hyperlink = OxmlElement("w:hyperlink")
    if anchor is not None:
        hyperlink.set(qn("w:anchor"), anchor)
    if url is not None:
        relationship_id = paragraph.part.relate_to(
            url,
            RELATIONSHIP_TYPE.HYPERLINK,
            is_external=True,
        )
        hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(TABLE_FONT_SIZE.pt * 2)))
    run_properties.append(size)

    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_external_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    append_hyperlink_run(paragraph, text=text, anchor=None, url=url)


def add_internal_hyperlink(paragraph: Paragraph, text: str, anchor: str) -> None:
    append_hyperlink_run(paragraph, text=text, anchor=anchor, url=None)


def render_markdownish_lines(document: DocxDocument, lines: list[str]) -> None:
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("- "):
            while index < len(lines) and lines[index].strip().startswith("- "):
                bullet = normalize_inline_markdown(lines[index].strip()[2:])
                document.add_paragraph(bullet, style="List Bullet")
                index += 1
            continue

        paragraph_lines = [normalize_inline_markdown(stripped)]
        index += 1
        while index < len(lines):
            next_stripped = lines[index].strip()
            if not next_stripped or next_stripped.startswith("- "):
                break
            paragraph_lines.append(normalize_inline_markdown(next_stripped))
            index += 1

        paragraph = document.add_paragraph(" ".join(paragraph_lines))
        for run in paragraph.runs:
            set_run_font(run)


def configure_document(document: DocxDocument) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = BODY_FONT_SIZE


def set_cell_text(cell: _Cell, text: str, *, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold, size=TABLE_FONT_SIZE)


def apply_table_grid_style(table: Table) -> None:
    table.style = "Table Grid"
    table.autofit = False


def set_table_widths(table: Table, widths_inches: list[float]) -> None:
    for row in table.rows:
        for index, width in enumerate(widths_inches):
            row.cells[index].width = Inches(width)


def add_report_table(
    document: DocxDocument,
    rows: list[ReportRow],
) -> None:
    table = document.add_table(rows=len(rows) + 1, cols=9)
    apply_table_grid_style(table)

    headers = [
        "Priority",
        "Company",
        "Role",
        "Posted",
        "Location / Work Model",
        "Fit",
        "URL",
        "Already Applied",
        "Contacts",
    ]
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)

    for row_index, row in enumerate(rows, start=1):
        cells = table.rows[row_index].cells
        set_cell_text(cells[0], row.priority)
        set_cell_text(cells[1], row.company)
        set_cell_text(cells[2], row.role)
        set_cell_text(cells[3], row.posted)
        set_cell_text(cells[4], row.location_work_model)
        set_cell_text(cells[5], row.fit)
        set_cell_text(cells[7], row.already_applied)

        job_link_paragraph = cells[6].paragraphs[0]
        job_link_paragraph.clear()
        add_external_hyperlink(job_link_paragraph, "Job page", row.url)

        details_paragraph = cells[8].paragraphs[0]
        details_paragraph.clear()
        add_internal_hyperlink(details_paragraph, "Open details", build_row_bookmark_name(row))

    set_table_widths(table, [0.5, 1.25, 2.05, 0.8, 2.2, 0.45, 0.8, 0.9, 0.8])


def add_heading_with_bookmark(
    document: DocxDocument,
    text: str,
    *,
    level: int,
    bookmark_name: str | None,
    bookmark_manager: BookmarkManager,
) -> Paragraph:
    paragraph = document.add_heading(text, level=level)
    if bookmark_name is not None:
        bookmark_manager.add_paragraph_bookmark(paragraph, bookmark_name)
    return paragraph


def add_job_summary_table(document: DocxDocument, row: ReportRow) -> None:
    table = document.add_table(rows=8, cols=2)
    apply_table_grid_style(table)

    facts = [
        ("Decision", row.decision),
        ("Priority", row.priority),
        ("Company", row.company),
        ("Role", row.role),
        ("Posted", row.posted),
        ("Location / Work Model", row.location_work_model),
        ("Fit", row.fit),
        ("Already Applied", row.already_applied),
    ]

    for index, (label, value) in enumerate(facts):
        set_cell_text(table.rows[index].cells[0], label, bold=True)
        set_cell_text(table.rows[index].cells[1], value)

    set_table_widths(table, [1.8, 7.5])
    link_paragraph = document.add_paragraph()
    link_label = link_paragraph.add_run("Official posting: ")
    set_run_font(link_label, bold=True)
    add_external_hyperlink(link_paragraph, "Job page", row.url)


def add_contacts_table(document: DocxDocument, company_section: NetworkingCompanySection) -> None:
    table = document.add_table(rows=len(company_section.contacts) + 1, cols=7)
    apply_table_grid_style(table)

    headers = [
        "Priority",
        "Name",
        "Current Role",
        "LinkedIn",
        "Connected On",
        "Match Source",
        "Why Reach Out",
    ]
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)

    for row_index, contact in enumerate(company_section.contacts, start=1):
        cells = table.rows[row_index].cells
        set_cell_text(cells[0], contact.priority)
        set_cell_text(cells[1], contact.name)
        set_cell_text(cells[2], contact.current_role)
        set_cell_text(cells[4], contact.connected_on)
        set_cell_text(cells[5], contact.match_source)
        set_cell_text(cells[6], contact.why_reach_out)

        linkedin_paragraph = cells[3].paragraphs[0]
        linkedin_paragraph.clear()
        add_external_hyperlink(linkedin_paragraph, "Profile", contact.linkedin_url)

    set_table_widths(table, [0.5, 1.35, 2.2, 0.9, 1.0, 1.1, 2.4])


def add_message_body(document: DocxDocument, body: str) -> None:
    blocks = [block.strip() for block in body.split("\n\n") if block.strip()]
    for block in blocks:
        paragraph = document.add_paragraph(normalize_inline_markdown(block))
        for run in paragraph.runs:
            set_run_font(run)


def add_detail_section(
    document: DocxDocument,
    row: ReportRow,
    company_section: NetworkingCompanySection | None,
    messages: dict[tuple[str, str], IntroductionMessage],
    bookmark_manager: BookmarkManager,
) -> None:
    document.add_page_break()
    add_heading_with_bookmark(
        document,
        f"{row.decision} {row.priority}: {row.company} - {row.role}",
        level=1,
        bookmark_name=build_row_bookmark_name(row),
        bookmark_manager=bookmark_manager,
    )

    back_link = document.add_paragraph()
    back_label = back_link.add_run("Back to summary: ")
    set_run_font(back_label, bold=True)
    add_internal_hyperlink(back_link, row.decision, SUMMARY_BOOKMARKS[row.decision])
    separator = back_link.add_run(" | ")
    set_run_font(separator)
    add_internal_hyperlink(back_link, "Top", SUMMARY_TOP_BOOKMARK)

    add_job_summary_table(document, row)

    add_heading_with_bookmark(
        document,
        "Who to Contact",
        level=2,
        bookmark_name=None,
        bookmark_manager=bookmark_manager,
    )

    if company_section is None:
        document.add_paragraph("No networking information is available for this company.")
    elif company_section.contacts:
        add_contacts_table(document, company_section)
    elif company_section.no_contacts_message is not None:
        document.add_paragraph(company_section.no_contacts_message)
    else:
        document.add_paragraph("No current-company LinkedIn contacts found.")

    add_heading_with_bookmark(
        document,
        "Suggested Contact Messages",
        level=2,
        bookmark_name=None,
        bookmark_manager=bookmark_manager,
    )

    if company_section is None or not company_section.contacts:
        document.add_paragraph("No draft contact message is available for this job.")
        return

    message_count = 0
    for contact in company_section.contacts:
        message = messages.get((row.company, contact.name))
        if message is None:
            continue
        message_count += 1
        add_heading_with_bookmark(
            document,
            contact.name,
            level=3,
            bookmark_name=None,
            bookmark_manager=bookmark_manager,
        )
        add_message_body(document, message.body)

    if message_count == 0:
        document.add_paragraph("No draft contact message is available for this job.")


def write_docx_report(report_path: Path, networking_path: Path | None = None, output_path: Path | None = None) -> Path:
    resolved_networking_path: Path
    resolved_output_path: Path
    if networking_path is None or output_path is None:
        default_networking_path, default_output_path = build_default_paths_docx(report_path)
        resolved_networking_path = networking_path or default_networking_path
        resolved_output_path = output_path or default_output_path
    else:
        resolved_networking_path = networking_path
        resolved_output_path = output_path

    report_markdown = report_path.read_text(encoding="utf-8")
    networking_markdown = resolved_networking_path.read_text(encoding="utf-8")

    cleaned_report = remove_networking_appendix(report_markdown)
    sections = parse_report_sections(cleaned_report)
    report_rows = parse_report_rows(report_markdown)
    if not report_rows:
        raise ValueError("No Apply or Maybe rows were found in the report.")

    company_sections, messages = parse_networking_markdown(networking_markdown)

    document = Document()
    configure_document(document)
    bookmark_manager = BookmarkManager()

    title = extract_title(
        report_markdown=report_markdown,
        fallback_stem=report_path.stem,
    )
    title_paragraph = document.add_paragraph()
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, bold=True, size=Pt(22))
    bookmark_manager.add_paragraph_bookmark(title_paragraph, SUMMARY_TOP_BOOKMARK)

    render_markdownish_lines(document, sections.preamble_lines)

    for decision in ("Apply", "Maybe"):
        decision_rows = [row for row in report_rows if row.decision == decision]
        add_heading_with_bookmark(
            document,
            decision,
            level=1,
            bookmark_name=SUMMARY_BOOKMARKS[decision],
            bookmark_manager=bookmark_manager,
        )
        add_report_table(document, decision_rows)

    for section_name in ("Skip", "Gap Questions for Candidate", "Follow-up Questions for Candidate", "Notes"):
        lines = sections.section_lines.get(section_name)
        if not lines:
            continue
        add_heading_with_bookmark(
            document,
            section_name,
            level=1,
            bookmark_name=None,
            bookmark_manager=bookmark_manager,
        )
        render_markdownish_lines(document, lines)

    for row in report_rows:
        add_detail_section(
            document=document,
            row=row,
            company_section=company_sections.get(row.company),
            messages=messages,
            bookmark_manager=bookmark_manager,
        )

    resolved_output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(resolved_output_path)
    return resolved_output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Render a daily job search markdown report to DOCX.")
    parser.add_argument("report_path", type=Path, help="Path to the daily job search markdown report.")
    parser.add_argument(
        "--networking-path",
        type=Path,
        default=None,
        help="Optional path to the networking matches markdown file.",
    )
    parser.add_argument(
        "--output-path",
        type=Path,
        default=None,
        help="Optional output path for the DOCX report.",
    )
    args = parser.parse_args()

    output_path = write_docx_report(
        report_path=args.report_path,
        networking_path=args.networking_path,
        output_path=args.output_path,
    )
    print(output_path)


if __name__ == "__main__":
    main()
