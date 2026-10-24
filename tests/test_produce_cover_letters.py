from __future__ import annotations

import base64
import json
import os
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path
from unittest import mock

from docx import Document
from docx.shared import Pt

from scripts.mark_applied_jobs import write_tracker_jobs_snapshot
from scripts.produce_cover_letters import produce_cover_letters


ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def write_source_resume(path: Path) -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("Candidate Name")
    table.cell(0, 1).paragraphs[0].add_run("candidate@example.com • LinkedIn • Portfolio")
    table.cell(0, 1).add_paragraph("555-0100 • Example City, ST")
    document.save(path)


def write_ready_queue(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "# 2026-07-24 - Resume-Ready Jobs",
                "",
                "| Priority | Company | Role | Decision | Template | Fit | URL |",
                "|---:|---|---|---|---|---:|---|",
                "| 1 | ReadyCo | Principal Engineer | APPLY | software_engineering | 86 | https://example.com/jobs/123 |",
                "",
            ]
        ),
        encoding="utf-8",
    )


def write_review(path: Path, *, authorized: bool) -> None:
    path.write_text(
        json.dumps(
            {
                "generated_on": "2026-07-24",
                "reviews": [
                    {
                        "company": "ReadyCo",
                        "title": "Principal Engineer",
                        "url": "https://example.com/jobs/123",
                        "should_generate_resume": authorized,
                        "already_applied": False,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )


def write_empty_snapshot(root: Path) -> Path:
    path = root / "2026-07-24-applied-jobs.json"
    write_tracker_jobs_snapshot(path, jobs=[], ssh_host="test-host")
    return path


def fake_renderer(path: Path, render_dir: Path) -> list[Path]:
    if not path.is_file():
        raise FileNotFoundError(path)
    render_dir.mkdir(parents=True, exist_ok=True)
    page = render_dir / "page-1.png"
    page.write_bytes(ONE_PIXEL_PNG)
    return [page]


class ProduceCoverLettersTests(unittest.TestCase):
    def setUp(self) -> None:
        self.evidence_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.evidence_dir.cleanup)
        evidence = Path(self.evidence_dir.name) / "evidence.json"
        evidence.write_text(json.dumps({"evidence": ["Delivered verified customer support improvements."]}))
        patcher = mock.patch.dict(os.environ, {"JOBAGENT_CANDIDATE_EVIDENCE": str(evidence)})
        patcher.start()
        self.addCleanup(patcher.stop)
        self.environment = mock.patch.dict(
            os.environ,
            {
                "JOBAGENT_SOURCE_RESUME_GLOB": "*_SOURCE_*.docx",
                "JOBAGENT_LINKEDIN_URL": "https://www.linkedin.com/in/candidate-profile",
                "JOBAGENT_PORTFOLIO_URL": "https://portfolio.example.com",
            },
        )
        self.environment.start()

    def tearDown(self) -> None:
        self.environment.stop()

    def test_generated_cover_letter_uses_compact_signature_contact_block(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            write_source_resume(source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx")
            ready_path = root / "2026-07-24-resume-ready-jobs.md"
            write_ready_queue(ready_path)
            review_path = root / "2026-07-24-workflow-review.json"
            write_review(review_path, authorized=True)

            produce_cover_letters(
                resume_ready_path=ready_path,
                workflow_review_path=review_path,
                applied_jobs_path=write_empty_snapshot(root),
                source_dir=source_dir,
                output_dir=root / "cover_letters",
                generated_on="2026-07-26",
                render_fn=fake_renderer,
            )

            cover_letter = root / "cover_letters/2026-07-24-ReadyCo-Principal-Engineer-Candidate-Name-Cover-Letter.docx"
            document = Document(cover_letter)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            with zipfile.ZipFile(cover_letter, "r") as archive:
                settings_root = ET.fromstring(archive.read("word/settings.xml"))
            compatibility_mode = next(
                setting
                for setting in settings_root.findall(
                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}compatSetting"
                )
                if setting.attrib.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name"
                )
                == "compatibilityMode"
            )

        self.assertEqual("July 26, 2026", paragraphs[0])
        signature_index = paragraphs.index("Sincerely,")
        self.assertEqual(
            [
                "Sincerely,",
                "Candidate Name",
                "candidate@example.com",
                "LinkedIn | Portfolio",
                "555-0100",
                "Example City, ST",
            ],
            paragraphs[signature_index : signature_index + 6],
        )
        self.assertTrue(document.paragraphs[signature_index + 1].runs[0].bold)
        self.assertEqual(Pt(9.5), document.paragraphs[signature_index + 2].runs[0].font.size)
        self.assertEqual(
            "15",
            compatibility_mode.attrib[
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
            ],
        )

    def test_generated_cover_letter_links_linkedin_and_portfolio_in_sender_header(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            write_source_resume(source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx")
            ready_path = root / "2026-07-24-resume-ready-jobs.md"
            write_ready_queue(ready_path)
            review_path = root / "2026-07-24-workflow-review.json"
            write_review(review_path, authorized=True)

            produce_cover_letters(
                resume_ready_path=ready_path,
                workflow_review_path=review_path,
                applied_jobs_path=write_empty_snapshot(root),
                source_dir=source_dir,
                output_dir=root / "cover_letters",
                generated_on="2026-07-26",
                render_fn=fake_renderer,
            )

            cover_letter = root / "cover_letters/2026-07-24-ReadyCo-Principal-Engineer-Candidate-Name-Cover-Letter.docx"
            with zipfile.ZipFile(cover_letter, "r") as archive:
                document_root = ET.fromstring(archive.read("word/document.xml"))
                relationships_root = ET.fromstring(archive.read("word/_rels/document.xml.rels"))

        targets = {
            relationship.attrib["Id"]: relationship.attrib["Target"]
            for relationship in relationships_root
            if relationship.attrib.get("TargetMode") == "External"
        }
        hyperlinks = {
            "".join(element.itertext()): targets[
                element.attrib[
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                ]
            ]
            for element in document_root.findall(
                ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
            )
        }
        self.assertEqual("https://www.linkedin.com/in/candidate-profile", hyperlinks["LinkedIn"])
        self.assertEqual("https://portfolio.example.com", hyperlinks["Portfolio"])

    def test_generates_only_strictly_authorized_cover_letter(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            write_source_resume(source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx")
            ready_path = root / "2026-07-24-resume-ready-jobs.md"
            write_ready_queue(ready_path)
            review_path = root / "2026-07-24-workflow-review.json"
            write_review(review_path, authorized=True)

            result = produce_cover_letters(
                resume_ready_path=ready_path,
                workflow_review_path=review_path,
                applied_jobs_path=write_empty_snapshot(root),
                source_dir=source_dir,
                output_dir=root / "cover_letters",
                generated_on="2026-07-26",
                render_fn=fake_renderer,
            )

            cover_letter = root / "cover_letters/2026-07-24-ReadyCo-Principal-Engineer-Candidate-Name-Cover-Letter.docx"
            self.assertTrue(cover_letter.is_file())
            text = "\n".join(paragraph.text for paragraph in Document(cover_letter).paragraphs)
            self.assertIn("July 26, 2026", text)
            self.assertIn("Principal Engineer", text)
            self.assertIn("Delivered verified customer support improvements.", text)
            self.assertNotIn("Oracle Fusion CPQ Cloud", text)
            self.assertIn(
                "I am applying for the Principal Engineer position at ReadyCo.",
                text,
            )
            self.assertIn(
                "I would welcome a conversation about how that experience could support ReadyCo.",
                text,
            )
            self.assertNotIn("I would welcome the opportunity to discuss", text)
            self.assertEqual("generated", result.records[0].status)
            manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
            self.assertEqual("2026-07-24", manifest["authorization_run_date"])
            self.assertEqual("candidate-confirmed-evidence", manifest["writing_voice"])
            self.assertEqual(
                "candidate-confirmed-evidence",
                manifest["records"][0]["writing_voice"],
            )
            self.assertEqual("validated", manifest["records"][0]["validation_status"])

    def test_rejects_job_not_authorized_by_strict_review(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            write_source_resume(source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx")
            ready_path = root / "2026-07-24-resume-ready-jobs.md"
            write_ready_queue(ready_path)
            review_path = root / "2026-07-24-workflow-review.json"
            write_review(review_path, authorized=False)

            with self.assertRaisesRegex(ValueError, "does not authorize"):
                produce_cover_letters(
                    resume_ready_path=ready_path,
                    workflow_review_path=review_path,
                    applied_jobs_path=write_empty_snapshot(root),
                    source_dir=source_dir,
                    output_dir=root / "cover_letters",
                    generated_on="2026-07-26",
                    render_fn=fake_renderer,
                )
            self.assertFalse((root / "cover_letters").exists())

    def test_existing_legacy_cover_letter_is_regenerated_in_approved_voice(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source_dir = root / "Source Docs"
            source_dir.mkdir()
            write_source_resume(source_dir / "Candidate_Name_SOURCE_Software_Engineering.docx")
            ready_path = root / "2026-07-24-resume-ready-jobs.md"
            write_ready_queue(ready_path)
            review_path = root / "2026-07-24-workflow-review.json"
            write_review(review_path, authorized=True)
            output_dir = root / "cover_letters"
            output_dir.mkdir()
            legacy_path = output_dir / "2026-07-24-ReadyCo-Principal-Engineer-Candidate-Name-Cover-Letter.docx"
            legacy = Document()
            legacy.add_paragraph("July 26, 2026")
            legacy.add_paragraph("ReadyCo")
            legacy.add_paragraph("Principal Engineer")
            legacy.add_paragraph("Oracle Fusion CPQ Cloud")
            legacy.add_paragraph(
                "I would welcome the opportunity to discuss how my experience could support ReadyCo."
            )
            legacy.save(legacy_path)

            result = produce_cover_letters(
                resume_ready_path=ready_path,
                workflow_review_path=review_path,
                applied_jobs_path=write_empty_snapshot(root),
                source_dir=source_dir,
                output_dir=output_dir,
                generated_on="2026-07-26",
                render_fn=fake_renderer,
            )

            text = "\n".join(paragraph.text for paragraph in Document(legacy_path).paragraphs)

        self.assertEqual("generated", result.records[0].status)
        self.assertIn(
            "I would welcome a conversation about how that experience could support ReadyCo.",
            text,
        )
        self.assertNotIn("I would welcome the opportunity to discuss", text)


if __name__ == "__main__":
    unittest.main()
